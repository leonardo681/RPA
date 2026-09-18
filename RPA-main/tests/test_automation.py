"""Integração local dos lotes. SMTP e WhatsApp são sempre substituídos por mocks."""

from dataclasses import replace
from datetime import date, timedelta
from pathlib import Path
import smtplib
from threading import Event

import pytest
from openpyxl import load_workbook

from techsolutions.config import Settings
from techsolutions.services.billing import Repository
from techsolutions.reports.excel import export_report, read_report
from techsolutions.reports.pdf import generate_all_pdfs
from techsolutions.automation.batch import prepare_batch, run_batch
from techsolutions.automation.emailer import DeliveryError, send_email
from techsolutions.automation.whatsapp import validate_accessible_link


@pytest.fixture
def billing_batch(tmp_path):
    settings = Settings(base_dir=tmp_path, data_dir=tmp_path,
                        database_path=tmp_path / "test.sqlite3", pdf_dir=tmp_path / "pdfs",
                        report_path=tmp_path / "faturas_relatorio.xlsx",
                        public_base_url="http://192.168.1.10:5000", admin_password="local-only",
                        secret_key="test-only", smtp_user="sender@example.com", smtp_password="test-password")
    repo = Repository(settings)
    repo.init_db()
    for index in range(1, 3):
        client_id = repo.save_client({"nome": f"Cliente Exemplo {index}", "email": f"cliente{index}@example.com",
                                     "telefone": f"1198765432{index}", "endereco": "Rua Exemplo, 10"})
        repo.create_invoice(client_id, f"{index}25.50", (date.today() + timedelta(days=index)).isoformat())
    generate_all_pdfs(settings, repo, log=lambda message: None)
    export_report(settings, repo)
    return settings, repo


def run_quiet(settings, repo, **kwargs):
    return run_batch(settings, repo, "email", log=lambda message: None, **kwargs)


def test_simulation_reads_excel_preserves_finance_and_history(billing_batch, monkeypatch):
    settings, repo = billing_batch
    monkeypatch.setattr("techsolutions.automation.batch.send_email", lambda *_: pytest.fail("Rede acionada em simulação"))
    summary = run_quiet(settings, repo)
    assert summary["simulados"] == 2
    assert all(row["status_email"] == "Simulado" and not row["data_envio_email"] for row in read_report(settings.report_path))
    assert all(invoice["status"] == "Pendente" for invoice in repo.list_invoices())
    assert all(attempt["simulado"] for attempt in repo.notifications())
    export_report(settings, repo)
    assert all(row["status_email"] == "Simulado" for row in read_report(settings.report_path))
    assert len(prepare_batch(settings, repo, "email")) == 2


def test_completed_attempt_is_not_sent_twice_even_after_export(billing_batch, monkeypatch):
    settings, repo = billing_batch
    sent = []
    def fake_send(_, item):
        sent.append((item["invoice_id"], item["recipient"], Path(item["pdf_path"]).name))
        return {"status": "Aceito pelo SMTP", "evidence": "Servidor falso de teste aceitou."}
    monkeypatch.setattr("techsolutions.automation.batch.send_email", fake_send)
    approval = prepare_batch(settings, repo, "email")
    summary = run_quiet(settings, repo, demo=False, approved_batch=approval)
    assert summary["sucessos"] == 2
    assert sent == [(1, "cliente1@example.com", "fatura_000001_cliente_000001.pdf"),
                    (2, "cliente2@example.com", "fatura_000002_cliente_000002.pdf")]
    export_report(settings, repo)
    assert prepare_batch(settings, repo, "email", retry_failed=True) == []
    assert all(row["data_envio_email"] for row in read_report(settings.report_path))
    assert all(invoice["status"] == "Pendente" for invoice in repo.list_invoices())


def test_real_send_requires_review_and_rejects_stale_snapshot(billing_batch, monkeypatch):
    settings, repo = billing_batch
    monkeypatch.setattr("techsolutions.automation.batch.send_email", lambda *_: pytest.fail("Não deve enviar"))
    with pytest.raises(ValueError, match="revisão"):
        run_quiet(settings, repo, demo=False)
    approval = prepare_batch(settings, repo, "email")
    repo.mark_paid(1)
    export_report(settings, repo)
    with pytest.raises(ValueError, match="mudou"):
        run_quiet(settings, repo, demo=False, approved_batch=approval)
    assert not repo.notifications()


@pytest.mark.parametrize("column,value", [(2, 2), (5, "outro@example.com"), (6, "0.01"), (10, "https://example.com/other")])
def test_tampered_report_blocks_before_transmission(billing_batch, column, value):
    settings, repo = billing_batch
    workbook = load_workbook(settings.report_path)
    workbook["Faturas"].cell(2, column, value)
    workbook.save(settings.report_path)
    workbook.close()
    with pytest.raises(ValueError, match="divergente"):
        prepare_batch(settings, repo, "email")
    assert not repo.notifications()


def test_failed_can_be_retried_but_uncertain_and_paid_cannot(billing_batch, monkeypatch):
    settings, repo = billing_batch
    def fake_send(_, item):
        raise DeliveryError("Falha de conexão de teste", uncertain=item["invoice_id"] == 2)
    monkeypatch.setattr("techsolutions.automation.batch.send_email", fake_send)
    approval = prepare_batch(settings, repo, "email")
    assert run_quiet(settings, repo, demo=False, approved_batch=approval)["falhas"] == 2
    assert prepare_batch(settings, repo, "email") == []
    assert [item["invoice_id"] for item in prepare_batch(settings, repo, "email", retry_failed=True)] == [1]
    assert repo.notification_state(2, "email")["status"] == "Não confirmado"
    assert (settings.data_dir / "erros.csv").is_file()
    repo.mark_paid(1)
    export_report(settings, repo)
    assert prepare_batch(settings, repo, "email", retry_failed=True) == []


def test_pdf_swap_is_recorded_and_never_sent(billing_batch, monkeypatch):
    settings, repo = billing_batch
    files = sorted(settings.pdf_dir.glob("*.pdf"))
    files[0].write_bytes(files[1].read_bytes())
    items = prepare_batch(settings, repo, "email")
    assert items[0]["validation_error"]
    calls = []
    monkeypatch.setattr("techsolutions.automation.batch.send_email", lambda _, item: calls.append(item["invoice_id"]) or {"status": "Aceito pelo SMTP"})
    result = run_quiet(settings, repo, demo=False, approved_batch=items)
    assert result["falhas"] == 1 and calls == [2]
    assert repo.notification_state(1, "email")["status"] == "Falha"


def test_missing_pdf_in_demo_does_not_create_real_failure(billing_batch):
    settings, repo = billing_batch
    Path(repo.get_invoice(1)["caminho_pdf"]).unlink()
    result = run_quiet(settings, repo)
    assert result["falhas"] == 1
    assert all(row["simulado"] for row in repo.notifications())
    assert repo.notification_state(1, "email")["status"] == "Simulado"


def test_stop_preserves_completed_record(billing_batch, monkeypatch):
    settings, repo = billing_batch
    stop = Event()
    def fake_send(*_):
        stop.set()
        return {"status": "Aceito pelo SMTP"}
    monkeypatch.setattr("techsolutions.automation.batch.send_email", fake_send)
    result = run_quiet(settings, repo, demo=False, approved_batch=prepare_batch(settings, repo, "email"), stop_event=stop)
    assert result["sucessos"] == 1
    assert repo.notification_state(1, "email")["status"] == "Aceito pelo SMTP"
    assert repo.notification_state(2, "email")["status"] == "Pendente"


def test_test_destination_is_explicit_and_invoice_attachment_stays_associated(billing_batch):
    settings, repo = billing_batch
    settings = replace(settings, test_email="authorized@example.com")
    item = prepare_batch(settings, repo, "email")[0]
    assert item["test_override"] is True
    assert item["recipient"] == "authorized@example.com"
    assert item["original_recipient"] == "cliente1@example.com"
    assert "cliente_000001.pdf" in item["pdf_path"]


def test_smtp_acceptance_and_matching_attachment(billing_batch, monkeypatch):
    settings, repo = billing_batch
    item = prepare_batch(settings, repo, "email")[0]
    seen = {}
    class FakeSMTP:
        def __init__(self, host, port, **kwargs):
            assert host == "smtp.gmail.com" and port == 465 and "context" in kwargs
        def login(self, user, password):
            assert user == settings.smtp_user
        def send_message(self, message, **kwargs):
            seen["message"] = message
            assert kwargs["to_addrs"] == [item["recipient"]]
            return {}
        def quit(self):
            raise OSError("Conexão fechou após aceitação; não invalida DATA")
        def close(self):
            pass
    monkeypatch.setattr("techsolutions.automation.emailer.smtplib.SMTP_SSL", FakeSMTP)
    assert send_email(settings, item)["status"] == "Aceito pelo SMTP"
    attachment = list(seen["message"].iter_attachments())[0]
    assert attachment.get_filename() == Path(item["pdf_path"]).name
    assert attachment.get_payload(decode=True) == Path(item["pdf_path"]).read_bytes()


@pytest.mark.parametrize("during_send", [False, True])
def test_smtp_connection_failure_distinguishes_uncertain_transmission(billing_batch, monkeypatch, during_send):
    settings, repo = billing_batch
    class FakeSMTP:
        def __init__(self, *args, **kwargs):
            pass
        def login(self, *args):
            if not during_send:
                raise smtplib.SMTPServerDisconnected("mock")
        def send_message(self, *args, **kwargs):
            raise smtplib.SMTPServerDisconnected("mock")
        def quit(self):
            pass
    monkeypatch.setattr("techsolutions.automation.emailer.smtplib.SMTP_SSL", FakeSMTP)
    with pytest.raises(DeliveryError) as captured:
        send_email(settings, prepare_batch(settings, repo, "email")[0])
    assert captured.value.uncertain is during_send


@pytest.mark.parametrize("link", ["http://localhost:5000/f/token", "http://127.0.0.1/f/token", "http://[::1]/f/token", "http://0.0.0.0/f/token"])
def test_whatsapp_rejects_links_unreachable_from_recipient(link):
    with pytest.raises(DeliveryError, match="localhost"):
        validate_accessible_link(link)
