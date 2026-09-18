"""Excel é entrada da cobrança; SQLite preserva o histórico e os vínculos."""
from decimal import Decimal
from pathlib import Path
from threading import RLock
from uuid import uuid4

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

from .pdf import invoice_link

COLUMNS = ["fatura_id", "cliente_id", "nome", "telefone", "email", "valor", "vencimento", "status_fatura", "caminho_pdf", "link_fatura", "status_whatsapp", "data_envio_whatsapp", "status_email", "data_envio_email", "erro_whatsapp", "erro_email"]
_LOCK = RLock()


def _values(settings, repo, invoice):
    whatsapp = repo.notification_state(invoice["id"], "whatsapp")
    email = repo.notification_state(invoice["id"], "email")
    return [invoice["id"], invoice["cliente_id"], invoice["nome"], invoice["telefone"], invoice["email"],
            Decimal(invoice["valor_centavos"])/100, invoice["data_vencimento"], invoice["situacao"], invoice["caminho_pdf"], invoice_link(settings, invoice),
            whatsapp["status"], whatsapp["data_envio"], email["status"], email["data_envio"], whatsapp["erro"], email["erro"]]


def _write_row(sheet, index, values):
    for col, value in enumerate(values, 1):
        cell = sheet.cell(index, col, value)
        # Strings de usuário nunca viram fórmulas do Excel.
        if isinstance(value, str):
            cell.data_type = "s"
        if col == 6:
            cell.number_format = '"R$" #,##0.00'
        elif col in (4,):
            cell.number_format = "@"


def _save_atomic(book, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.stem + f".{uuid4().hex}.tmp.xlsx")
    try:
        book.save(temp)
        temp.replace(path)
    except PermissionError as exc:
        raise PermissionError("Feche a planilha no Excel e tente novamente. Resultados já gravados no banco são preservados; exporte novamente para sincronizar.") from exc
    finally:
        temp.unlink(missing_ok=True)
        book.close()


def export_report(settings, repo):
    with _LOCK:
        book = Workbook()
        sheet = book.active
        sheet.title = "Faturas"
        sheet.append(COLUMNS)
        for index, invoice in enumerate(repo.list_invoices(), 2):
            _write_row(sheet, index, _values(settings, repo, invoice))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163B61")
            cell.alignment = Alignment(wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(16, len(str(col[0].value))+3))
        if sheet.max_row > 1:
            table = Table(displayName="Cobrancas", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        _save_atomic(book, settings.report_path)
    return settings.report_path


def read_report(path):
    path = Path(path)
    if not path.is_file():
        raise ValueError("Relatório inexistente. Exporte o relatório Excel primeiro.")
    frame = pd.read_excel(path, sheet_name="Faturas", dtype=str, keep_default_na=False, engine="openpyxl")
    missing = set(COLUMNS)-set(frame.columns)
    if missing:
        raise ValueError("Colunas ausentes: " + ", ".join(sorted(missing)))
    rows = frame.to_dict("records")
    seen = set()
    for row in rows:
        try:
            row["fatura_id"] = int(row["fatura_id"])
            row["cliente_id"] = int(row["cliente_id"])
            amount = Decimal(row["valor"])
            if not amount.is_finite() or amount != amount.quantize(Decimal("0.01")):
                raise ValueError("Precisão monetária inválida no relatório.")
            row["valor"] = str(amount.quantize(Decimal("0.01")))
        except (ValueError, ArithmeticError) as exc:
            raise ValueError("Relatório com identificador ou valor inválido.") from exc
        if row["fatura_id"] in seen:
            raise ValueError("Relatório contém fatura_id duplicado; exporte novamente.")
        seen.add(row["fatura_id"])
    return rows


def sync_report_row(settings, repo, invoice_id, path=None):
    report = Path(path or settings.report_path)
    with _LOCK:
        book = load_workbook(report)
        sheet = book["Faturas"]
        headers = [cell.value for cell in sheet[1]]
        if headers != COLUMNS:
            book.close()
            raise ValueError("Estrutura do relatório alterada. Exporte novamente.")
        matches = [row[0].row for row in sheet.iter_rows(min_row=2) if str(row[0].value) == str(invoice_id)]
        if len(matches) != 1:
            book.close()
            raise ValueError("Fatura ausente ou duplicada na planilha. Exporte novamente.")
        _write_row(sheet, matches[0], _values(settings, repo, repo.get_invoice(invoice_id)))
        _save_atomic(book, report)
